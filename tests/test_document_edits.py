from docx import Document
from meeting_notes.summary_models import SummaryProject,parse_summary,save_summary,load_summary
from meeting_notes.summary_export import export_summary,preview_html
from meeting_notes.summary_layout import document_blocks,static_text

def test_all_document_text_edits_roundtrip_and_export(tmp_path):
    p=SummaryProject(items=parse_summary('- 第一條。\n- 不要這一條。'))
    p.items[1].included=False
    p.document_edits={'company':'測試公司','meeting_name':'可改的標題',
        'meeting_time':'會議時間：使用者填寫', 'heading_chair':'主席裁示（已編輯）',
        'heading_reports':'工作報告', 'recorder':'紀錄：測試人員', 'tools':'使用工具：自行填寫'}
    p.items[0].text='改寫第一段。\n另加一段。'
    save_summary(p,tmp_path/'saved.summary.json')
    p=load_summary(tmp_path/'saved.summary.json')
    export_summary(p,tmp_path/'out.docx')
    d=Document(tmp_path/'out.docx');text='\n'.join(x.text for x in d.paragraphs)
    header='\n'.join(x.text for x in d.sections[0].header.paragraphs)
    company=''.join(c.text for t in d.sections[0].header.tables for row in t.rows for c in row.cells)
    assert '測試公司' in company and '可改的標題' in header
    for value in ['會議時間：使用者填寫','主席裁示（已編輯）','工作報告','紀錄：測試人員','自行填寫','另加一段。']:
        assert value in text and value in preview_html(p)
    assert '不要這一條' not in text

def test_blank_override_is_kept_and_hidden_items_do_not_erase_text():
    p=SummaryProject(items=parse_summary('- 保留修改。'))
    p.document_edits={'meeting_name':''}
    assert static_text(p,'meeting_name')==''
    item=p.items[0];item.included=False
    assert not any(b.key==item.id for b in document_blocks(p))
    item.included=True
    assert next(b.text for b in document_blocks(p) if b.key==item.id)=='保留修改。'
