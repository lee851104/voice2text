import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import pytest
from docx import Document
from PySide6.QtCore import Qt, QPointF, QPoint, QMimeData
from PySide6.QtGui import QDragEnterEvent, QDropEvent
from PySide6.QtWidgets import QApplication, QComboBox
from meeting_notes.summary_ui import SummaryWindow
from meeting_notes.summary_models import SummaryProject, parse_summary, save_summary, load_summary
from meeting_notes.summary_export import export_summary
from meeting_notes.summary_transfer import block_mime

pytestmark=pytest.mark.filterwarnings('ignore:Function.*QDragEnterEvent:DeprecationWarning')

@pytest.fixture(scope='module')
def app(): return QApplication.instance() or QApplication([])

@pytest.fixture
def window(app):
    w=SummaryWindow();w.resize(1366,768);w.show();app.processEvents()
    w.source_input.setPlainText('- 第一則原文。\n- 第二則原文。\n- 第三則原文。');w.go_review();app.processEvents()
    yield w
    w.dirty=False;w.close();app.processEvents()

def test_new_source_starts_included_and_arrow_readds_without_duplicates(window):
    w=window
    assert len(w.project.selected())==3
    assert all(w.table.item(row,0).checkState()==Qt.CheckState.Checked for row in range(3))
    assert not w.findChildren(QComboBox)
    item=w.project.items[0]
    w.remove_all()
    w.table.selectRow(0);w.add_selected_button.click()
    assert item.included and item.id in w.preview.edits
    w.add_selected_button.click()
    assert len(w.project.selected())==1
    w.add_all_button.click()
    assert len(w.project.selected())==3

def test_remove_readd_keeps_edits_and_original(window):
    w=window;item=w.project.items[0]
    w.insert_items([item.id],'chair')
    w.preview.edits[item.id].setPlainText('手動修改後。')
    w.preview.remove_buttons[item.id].click()
    assert not item.included and item.id not in w.preview.edits
    w.insert_items([item.id],'motions')
    assert w.preview.edits[item.id].toPlainText()=='手動修改後。'
    assert item.source=='第一則原文。' and item.section=='motions'

def test_move_order_export_and_project_roundtrip(window,tmp_path):
    w=window;a,b,c=w.project.items
    w.insert_items([a.id,b.id,c.id],'chair')
    w.insert_items([c.id],'chair',a.id)
    w.insert_items([b.id],'motions')
    assert [i.id for i in w.project.selected() if i.section=='chair']==[c.id,a.id]
    assert b.section=='motions'
    w.remove_items([a.id])
    w.preview.edits['meeting_time'].setPlainText('會議時間：測試時間')
    path=tmp_path/'drag.summary.json';save_summary(w.project,path)
    assert load_summary(path)==w.project
    output=tmp_path/'drag.docx';export_summary(w.project,output)
    text='\n'.join(p.text for p in Document(output).paragraphs)
    assert a.text not in text and text.index(c.text)<text.index(b.text)
    assert '測試時間' in text

def test_native_drop_on_heading_adds_whole_block_without_replacing_title(window,app):
    w=window;item=w.project.items[0]
    w.table.selectRow(0)
    mime=w.table.mimeData(w.table.selectedItems())
    target=w.preview.edits['heading_chair'].viewport()
    position=QPoint(10,10)
    enter=QDragEnterEvent(position,Qt.DropAction.CopyAction,mime,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier)
    QApplication.sendEvent(target,enter)
    assert enter.isAccepted()
    drop=QDropEvent(QPointF(position),Qt.DropAction.CopyAction,mime,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier)
    QApplication.sendEvent(target,drop);app.processEvents()
    assert drop.isAccepted() and item.included and item.section=='chair'
    assert w.preview.edits['heading_chair'].toPlainText()=='主席宣達事項'
    assert w.preview.edits[item.id].toPlainText()==item.text

def test_invalid_drop_and_old_project_remain_safe(window,app):
    w=window
    w.remove_all()
    w.insert_items(['unknown'],'chair')
    assert not w.project.selected()
    old=SummaryProject(items=parse_summary('- 已儲存內容。'))
    w.set_project(old)
    assert old.items[0].id in w.preview.edits
    mime=QMimeData();mime.setText('外部純文字')
    target=w.preview.edits['heading_chair'].viewport()
    enter=QDragEnterEvent(QPoint(10,10),Qt.DropAction.CopyAction,mime,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier)
    QApplication.sendEvent(target,enter)
    assert not enter.isAccepted()

def test_native_document_reorder_and_drag_back_to_source(window,app):
    w=window;a,b,c=w.project.items
    w.add_all()
    mime=block_mime([c.id],w.session_token)
    target=w.preview.edits[a.id].viewport()
    enter=QDragEnterEvent(QPoint(8,8),Qt.DropAction.CopyAction,mime,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier)
    QApplication.sendEvent(target,enter)
    drop=QDropEvent(QPointF(8,8),Qt.DropAction.CopyAction,mime,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier)
    QApplication.sendEvent(target,drop);app.processEvents()
    assert [i.id for i in w.project.selected()]==[c.id,a.id,b.id]
    target=w.table.viewport()
    enter=QDragEnterEvent(QPoint(8,8),Qt.DropAction.CopyAction,mime,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier)
    QApplication.sendEvent(target,enter)
    drop=QDropEvent(QPointF(8,8),Qt.DropAction.CopyAction,mime,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier)
    QApplication.sendEvent(target,drop);app.processEvents()
    assert drop.isAccepted() and not c.included
    assert c.id not in w.preview.edits

def test_multiselect_arrow_and_click_destination(window,app):
    w=window;a,b,c=w.project.items
    w.remove_all()
    w.preview.edits['heading_motions'].setFocus();app.processEvents()
    w.table.item(0,1).setSelected(True);w.table.item(2,1).setSelected(True)
    w.add_selected_button.click()
    assert a.included and c.included and not b.included
    assert a.section==c.section=='motions'
    assert w.table.item(0,1).data(Qt.ItemDataRole.UserRole)==a.id
    assert w.table.item(0,1).text()==a.source

def test_heading_drop_can_move_existing_paragraph_to_end(window):
    w=window;a,b,c=w.project.items;w.add_all()
    w.insert_items([a.id],'chair')
    assert [i.id for i in w.project.selected()]==[b.id,c.id,a.id]

def test_long_source_paragraph_can_scroll_through_middle(window,app):
    w=window
    w.source_input.setPlainText('- '+('這是很長的文字內容。'*600))
    w.project.items=[];w.go_review();app.processEvents()
    bar=w.table.verticalScrollBar()
    assert w.table.rowHeight(0)>w.table.viewport().height()*2
    assert bar.maximum()>w.table.viewport().height()
    bar.setValue(900)
    assert -1100<w.table.rowViewportPosition(0)<-700

def test_add_all_uses_source_order_after_removal(window):
    w=window;a,b,c=w.project.items
    w.remove_all()
    w.insert_items([a.id],'chair');w.remove_items([a.id]);w.add_all()
    assert [i.id for i in w.project.selected()]==[a.id,b.id,c.id]

def test_saved_source_order_survives_document_reorder(window,tmp_path):
    w=window;a,b,c=w.project.items;w.add_all()
    w.insert_items([c.id],'reports',a.id)
    path=tmp_path/'order.summary.json';save_summary(w.project,path)
    w.set_project(load_summary(path))
    assert [w.table.item(i,1).data(Qt.ItemDataRole.UserRole) for i in range(3)]==[a.id,b.id,c.id]
    assert [i.id for i in w.project.selected()]==[c.id,a.id,b.id]

def test_checkbox_tracks_id_after_drag_and_keeps_edits_position(window,tmp_path):
    w=window;a,b,c=w.project.items
    w.insert_items([c.id],'motions')
    w.insert_items([b.id],'reports',a.id)
    w.preview.edits[a.id].setPlainText('修訂後的第一段。')
    w.table.item(0,0).setCheckState(Qt.CheckState.Unchecked)
    assert not a.included and b.included and c.included
    assert a.id not in w.preview.edits
    assert '修訂後的第一段' not in w.preview.toPlainText()
    w.table.item(0,0).setCheckState(Qt.CheckState.Checked)
    assert a.section=='reports' and a.text=='修訂後的第一段。'
    assert [i.id for i in w.project.items if i.section=='reports']==[b.id,a.id]
    w.preview.remove_buttons[c.id].click()
    assert w.table.item(2,0).checkState()==Qt.CheckState.Unchecked
    path=tmp_path/'checks.summary.json';save_summary(w.project,path)
    w.set_project(load_summary(path))
    assert w.table.item(2,0).checkState()==Qt.CheckState.Unchecked
    output=tmp_path/'checks.docx';export_summary(w.project,output)
    text='\n'.join(p.text for p in Document(output).paragraphs)
    assert c.text not in text and a.text in text and text.index(b.text)<text.index(a.text)

def test_select_all_preserves_sections_and_does_not_change_source(window):
    w=window;a,b,c=w.project.items
    w.insert_items([b.id],'motions')
    w.select_none_button.click()
    assert not w.project.selected() and not w.export_button.isEnabled()
    w.select_all_button.click()
    assert len(w.project.selected())==3 and b.section=='motions'
    assert all(w.table.item(row,0).checkState()==Qt.CheckState.Checked for row in range(3))
    assert a.source=='第一則原文。'

def test_new_blank_checkbox_is_visible_and_can_be_clicked_in_long_list(window,app):
    from PySide6.QtTest import QTest
    w=window
    w.source_input.setPlainText('\n'.join('- 測試段落'+str(n) for n in range(30)))
    w.project.items=[];w.go_review();QTest.qWait(30)
    w.add_item();QTest.qWait(30)
    item=w.project.items[-1];row=w.table.rowCount()-1
    rect=w.table.visualItemRect(w.table.item(row,0))
    point=QPoint(rect.left()+15,rect.center().y())
    assert w.table.viewport().rect().contains(rect), '新增段落的勾選框應完整顯示，不應落在清單底部之外'
    QTest.mouseClick(w.table.viewport(),Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier,point)
    app.processEvents()
    assert not item.included and item.id not in w.preview.edits
    QTest.mouseClick(w.table.viewport(),Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier,point)
    app.processEvents()
    assert item.included and item.id in w.preview.edits

@pytest.mark.parametrize('included',[False,True])
def test_delete_manual_blank_removes_left_row_and_saved_data(window,tmp_path,included):
    w=window;original_ids=[i.id for i in w.project.items]
    w.add_item();manual=w.project.items[-1]
    if not included:w.remove_items([manual.id])
    row=w.table.rowCount()-1
    button=w.table.cellWidget(row,2)
    assert button is not None,'手動段落需要左側刪除按鈕，即使未勾選也可刪除'
    button.click()
    assert [i.id for i in w.project.items]==original_ids
    assert w.table.rowCount()==3 and manual.id not in w.preview.edits
    assert manual.id not in w.project.source_order
    path=tmp_path/'deleted.summary.json';save_summary(w.project,path)
    w.set_project(load_summary(path));w.set_all(True)
    assert all(i.id!=manual.id for i in w.project.items)
    assert w.table.isColumnHidden(2)

def test_deleting_written_manual_item_can_cancel_and_does_not_undo_restore(window,monkeypatch):
    w=window;w.add_item();manual=w.project.items[-1]
    w.preview.edits[manual.id].setPlainText('保留這份人工紀錄。')
    w.apply_rewrite({manual.id:'改寫後的人工紀錄。'})
    monkeypatch.setattr(w,'confirm',lambda *args:False)
    w.table.cellWidget(w.table.rowCount()-1,2).click()
    assert any(i.id==manual.id for i in w.project.items)
    monkeypatch.setattr(w,'confirm',lambda *args:True)
    w.table.cellWidget(w.table.rowCount()-1,2).click()
    w.undo_rewrite();w.set_all(True)
    assert all(i.id!=manual.id for i in w.project.items)
    assert manual.id not in w.preview.edits

def test_source_paragraph_has_no_manual_delete_action(window):
    w=window;original=w.project.model_copy(deep=True)
    assert w.table.cellWidget(0,2) is None
    w.delete_manual_item(w.project.items[0].id)
    assert w.project==original
