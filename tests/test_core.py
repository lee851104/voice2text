import json
from pathlib import Path

import pytest

from meeting_notes.models import Project, Minutes, Evidence, Action
from meeting_notes.storage import save_project, load_project
from meeting_notes.export import export_docx


def fixture_project():
    project = Project(title="產品進度會議", date="2026-10-12")
    project.add_chunk({"segments": [
        {"speaker": "A", "start": 0, "end": 3, "text": "我會整理測試報告。"},
        {"speaker": "B", "start": 4, "end": 7, "text": "期限先不決定。"},
    ]}, 0, 0)
    return project


def test_chunk_labels_remain_distinct_and_timestamps_are_absolute():
    p = fixture_project()
    p.add_chunk({"segments": [{"speaker": "A", "start": 2, "end": 5, "text": "下一段"}]}, 1, 5400)
    assert p.segments[0].speaker != p.segments[-1].speaker
    assert p.segments[-1].start == 5402
    assert p.segments[-1].end == 5405
    assert len(p.speakers) == 3


def test_rename_merge_and_edit_invalidate_summary_without_altering_raw_text():
    p = fixture_project()
    first, second = list(p.speakers)
    p.minutes = Minutes(summary=[Evidence(text="已有安排", sources=["s0-0"])])
    p.rename_speaker(first, "王小姐")
    assert p.minutes is None
    p.reassign_segment("s0-1", first)
    p.edit_segment("s0-0", "我會整理正式報告。")
    assert p.segments[1].speaker == first
    assert p.segments[0].original_text == "我會整理測試報告。"
    assert p.segments[0].text == "我會整理正式報告。"
    assert p.speakers[first] == "王小姐"


def test_project_roundtrip_does_not_accept_or_store_key(tmp_path):
    p = fixture_project()
    target = tmp_path / "會議.meeting.json"
    save_project(p, target)
    loaded = load_project(target)
    assert loaded.segments == p.segments
    content = json.loads(target.read_text(encoding="utf-8"))
    assert "api_key" not in content
    content["api_key"] = "a-secret-that-must-not-be-loaded"
    target.write_text(json.dumps(content), encoding="utf-8")
    with pytest.raises(ValueError):
        load_project(target)


def test_word_uses_corrected_names_full_transcript_and_unknown_deadlines(tmp_path):
    from docx import Document
    p = fixture_project()
    p.rename_speaker(p.segments[0].speaker, "王小姐")
    p.minutes = Minutes(summary=[Evidence(text="討論測試報告", sources=["s0-0"])],
        actions=[Action(task="整理測試報告", owner="王小姐", deadline="未指定", sources=["s0-0"])])
    target = tmp_path / "會議.docx"
    export_docx(p, target)
    doc = Document(target)
    text = "\n".join(x.text for x in doc.paragraphs)
    tables = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert "王小姐" in text
    assert "00:00:04" in text
    assert "期限先不決定。" in text
    assert "未指定" in tables
    assert "討論測試報告" in text
    assert doc.core_properties.author == ""


def test_reject_invalid_segments_instead_of_partial_mutation():
    p = fixture_project()
    with pytest.raises(ValueError):
        p.add_chunk({"segments": [{"speaker":"C", "start":8,"end":2,"text":"bad"}]}, 1, 0)
    assert len(p.segments) == 2


@pytest.mark.parametrize("value", ["../../target", "D:/outside", "f" * 63, "z" * 64])
def test_imported_project_cannot_supply_path_as_fingerprint(value):
    with pytest.raises(ValueError):
        Project(audio_hash=value)


def test_unhashed_projects_have_independent_safe_session_ids():
    one, two = Project(), Project()
    assert one.session_id != two.session_id
    assert len(one.session_id) == 32
    with pytest.raises(ValueError):
        Project(session_id="../../elsewhere")


def test_minutes_schema_has_no_default_keywords_for_strict_api_contract():
    def walk(schema):
        if isinstance(schema, dict):
            assert "default" not in schema
            for value in schema.values():
                walk(value)
        elif isinstance(schema, list):
            for value in schema:
                walk(value)
    walk(Minutes.model_json_schema())
