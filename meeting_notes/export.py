from __future__ import annotations

import os
from pathlib import Path
import re
import tempfile

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .models import Project, timestamp


def clean(text: str) -> str:
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", text)


def export_docx(project: Project, destination: Path):
    if not project.segments:
        raise ValueError("尚無可匯出的逐字稿")
    if project.minutes:
        project.validate_minutes(project.minutes)
    doc = Document()
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = clean(project.title)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet"):
        style = doc.styles[name]
        style.font.name = "Microsoft JhengHei"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.font.color.rgb = RGBColor.from_string("182536")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Title"].font.size = Pt(24)
    for border in doc.styles["Title"].element.findall(".//" + qn("w:pBdr")):
        border.getparent().remove(border)
    for name, size, before in (("Heading 1", 14, 12), ("Heading 2", 11.5, 8)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(4)
    doc.add_paragraph(clean(project.title), "Title")
    doc.add_paragraph(f"會議日期  {clean(project.date) or '未指定'}")
    used = dict.fromkeys(s.speaker for s in project.segments)
    doc.add_paragraph("參與者  " + "、".join(clean(project.speakers[s]) for s in used))
    if not project.transcription_complete:
        doc.add_paragraph("此份為部分逐字稿，錄音尚未全部轉錄完成。")
    doc.add_paragraph("本紀錄依錄音辨識與人工校對內容整理；請核對姓名、決議與待辦事項。")
    by_id = {s.id: s for s in project.segments}

    def source_text(ids):
        return "、".join(dict.fromkeys(timestamp(by_id[x].start) for x in ids))

    def bullet(entry):
        paragraph = doc.add_paragraph(clean(entry.text), "List Bullet")
        run = paragraph.add_run(f"  ［{source_text(entry.sources)}］")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("627084")

    minutes = project.minutes
    if minutes:
        doc.add_heading("會議摘要", 1)
        for entry in minutes.summary:
            bullet(entry)
        if not minutes.summary:
            doc.add_paragraph("未整理出可確認的摘要。")
        doc.add_heading("議題與討論", 1)
        for topic in minutes.topics:
            doc.add_heading(clean(topic.title), 2)
            for entry in topic.points:
                bullet(entry)
        doc.add_heading("各人發言重點", 1)
        people = {p.speaker: p for p in minutes.people}
        for speaker in used:
            doc.add_heading(clean(project.speakers[speaker]), 2)
            if speaker in people:
                for entry in people[speaker].points:
                    bullet(entry)
            else:
                doc.add_paragraph("請參閱附錄中此人的完整發言。")
        doc.add_heading("會議決議", 1)
        for entry in minutes.decisions:
            bullet(entry)
        if not minutes.decisions:
            doc.add_paragraph("未記錄到明確決議。")
        doc.add_heading("待辦事項", 1)
        if minutes.actions:
            table = doc.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            table.autofit = False
            widths = (7.6, 2.6, 3, 3.4)
            for column, width in zip(table.columns, widths):
                column.width = Cm(width)
            for cell, title in zip(table.rows[0].cells, ("工作內容", "負責人", "期限", "錄音時間")):
                cell.text = title
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for action in minutes.actions:
                row = table.add_row()
                for cell, value in zip(row.cells, (action.task, action.owner or "未指定", action.deadline or "未指定", source_text(action.sources))):
                    cell.text = clean(value)
            for row_index, row in enumerate(table.rows):
                no_split = OxmlElement("w:cantSplit")
                row._tr.get_or_add_trPr().append(no_split)
                for cell, width in zip(row.cells, widths):
                    cell.width = Cm(width)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(3)
                        if row_index == 0:
                            paragraph.paragraph_format.keep_with_next = True
                            for run in paragraph.runs:
                                run.bold = True
        else:
            doc.add_paragraph("未記錄到明確待辦事項。")
    else:
        doc.add_paragraph("此份僅含逐字稿；尚未產生會議摘要。")
    doc.add_page_break()
    doc.add_heading("附錄 完整逐字稿", 1)
    for segment in project.segments:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.keep_with_next = True
        label = f"{timestamp(segment.start)} – {timestamp(segment.end)}　{project.speakers[segment.speaker]}"
        paragraph.add_run(clean(label)).bold = True
        doc.add_paragraph(clean(segment.text))
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("會議紀錄　｜　")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(suffix=".docx", dir=destination.parent)
    os.close(fd)
    try:
        doc.save(temporary)
        os.replace(temporary, destination)
    finally:
        Path(temporary).unlink(missing_ok=True)
