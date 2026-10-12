"""以既定範本匯出；畫面提供相同章節的內容預覽。"""
from copy import deepcopy
from html import escape
from pathlib import Path
import sys
import os
import tempfile
from docx import Document
from docx.text.paragraph import Paragraph
from .summary_models import SummaryProject, SECTIONS
from .language import traditional
from .summary_layout import TEMPLATE_SLOTS,static_text,document_blocks

def template_path():
    root=Path(getattr(sys,'_MEIPASS',Path(__file__).resolve().parents[1]))
    return root/'assets'/'department-template.docx'

def replace_text(paragraph, text):
    props=deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    paragraph.clear()
    run=paragraph.add_run(traditional(text))
    if props is not None: run._r.insert(0,props)

def grouped(project):
    groups={key:[] for key in SECTIONS}
    for item in project.selected(): groups[item.section].append(traditional(item.text.strip()))
    return groups

def export_summary(project:SummaryProject, path):
    if not project.selected(): raise ValueError('請至少勾選一則有內容的條目。')
    doc=Document(template_path())
    groups=grouped(project)
    for p in list(doc.paragraphs):
        key=p.text.removeprefix('{{').removesuffix('}}')
        if key in groups and p.text=='{{'+key+'}}':
            for value in groups[key]:
                clone=deepcopy(p._p)
                p._p.addprevious(clone)
                replace_text(Paragraph(clone,p._parent),value)
            p._p.getparent().remove(p._p)
        else:
            slot=TEMPLATE_SLOTS.get(p.text)
            if slot: replace_text(p,static_text(project,slot))
    for section in doc.sections:
        for p in section.header.paragraphs:
            if '{{meeting_name}}' in p.text: replace_text(p,static_text(project,'meeting_name'))
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if p.text=='桃園大眾捷運股份有限公司':replace_text(p,static_text(project,'company'))
    path=Path(path)
    descriptor,temporary=tempfile.mkstemp(prefix=path.name+'.',suffix='.tmp',dir=path.parent)
    os.close(descriptor)
    try:
        doc.save(temporary)
        os.replace(temporary,path)
    finally:
        if Path(temporary).exists():Path(temporary).unlink()

def preview_html(project:SummaryProject):
    e=lambda value:escape(traditional(value)).replace('\n','<br>')
    parts=['<html><body style="font-family:標楷體; font-size:14pt;">']
    for block in document_blocks(project):
        tag='h2' if block.kind=='company' else 'h3' if block.kind=='title' else 'p'
        align='center' if block.kind in ('company','title') else 'right' if block.kind=='recorder' else 'left'
        parts.append(f'<{tag} align="{align}">'+e(block.prefix+block.text)+f'</{tag}>')
    return ''.join(parts)+'</body></html>'
